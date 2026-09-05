import io
import unittest
import zipfile
import subprocess
import sys
from pathlib import Path
import extractor
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def archive(files):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as bundle:
        for name, content in files.items():
            bundle.writestr(name, content)
    return stream.getvalue()


class ExtractionTests(unittest.TestCase):
    def test_docx_paragraphs_and_table_cells_keep_locations(self):
        raw = archive({"word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>登录失败</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>重现步骤</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>'})
        result = extractor.extract(raw, "docx")
        self.assertEqual([r["text"] for r in result["records"]], ["登录失败", "重现步骤"])
        self.assertIn("paragraph:2", result["records"][1]["location"])
        self.assertIn("table:1:row:1:cell:1", result["records"][1]["location"])
        self.assertFalse(result["truncated"])

    def test_xlsx_all_sheets_and_formulas_are_data_not_executed(self):
        workbook = Workbook()
        workbook.active.title = "缺陷"
        workbook.active.append(["编号", "现象"])
        workbook.active.append([1, "登录失败"])
        second = workbook.create_sheet("补充")
        second["B2"] = "=HYPERLINK(\"https://invalid.example\",\"do not fetch\")"
        stream = io.BytesIO()
        workbook.save(stream)
        result = extractor.extract(stream.getvalue(), "xlsx")
        self.assertTrue(any("缺陷!A2" in r["location"] for r in result["records"]))
        self.assertTrue(any("HYPERLINK" in r["text"] for r in result["records"]))
        self.assertIn("formulas_not_evaluated", result["warnings"])
        self.assertTrue(result["truncated"])

    def test_pdf_text_keeps_page_and_empty_pages_are_incomplete(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 10 100 Td (Login failed) Tj ET")
        page[NameObject("/Contents")] = content
        writer.add_blank_page(width=200, height=200)
        stream = io.BytesIO()
        writer.write(stream)
        result = extractor.extract(stream.getvalue(), "pdf")
        self.assertIn("Login failed", result["records"][0]["text"])
        self.assertEqual(result["records"][0]["location"], "page:1")
        self.assertTrue(result["truncated"])

    def test_rejects_archive_traversal_macros_entities_and_bombs(self):
        for files in [{"../escape": "bad"}, {"word/vbaProject.bin": "macro"},
                      {"word/document.xml": '<!DOCTYPE doc [<!ENTITY x SYSTEM "file:///etc/passwd">]><doc>&x;</doc>'},
                      {"word/document.xml": "a" * 2_000_000}]:
            with self.assertRaises(Exception):
                extractor.extract(archive(files), "docx")

    def test_output_limit_is_explicit_not_silent(self):
        previous = extractor.MAX_CHARACTERS
        extractor.MAX_CHARACTERS = 5
        try:
            result = extractor.extract(archive({"word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>abcdef</w:t></w:r></w:p></w:document>'}), "docx")
            self.assertTrue(result["truncated"])
            self.assertIn("output_truncated", result["warnings"])
        finally:
            extractor.MAX_CHARACTERS = previous

    def test_pdf_indirect_resources_are_supported(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        page[NameObject("/Resources")] = writer._add_object(DictionaryObject())
        stream = io.BytesIO()
        writer.write(stream)
        result = extractor.extract(stream.getvalue(), "pdf")
        self.assertIn("unread_empty_or_scanned_page", result["warnings"])

    def test_encrypted_pdf_is_rejected(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.encrypt("fixture-password")
        stream = io.BytesIO()
        writer.write(stream)
        with self.assertRaises(ValueError):
            extractor.extract(stream.getvalue(), "pdf")

    def test_docx_fields_and_external_links_remain_incomplete(self):
        raw = archive({"word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:fldSimple w:instr="INCLUDETEXT file:///etc/passwd"><w:r><w:t>cached result</w:t></w:r></w:fldSimple></w:p></w:document>'})
        result = extractor.extract(raw, "docx")
        self.assertIn("unread_revision_or_field", result["warnings"])

    def test_external_relationships_are_not_followed(self):
        raw = archive({"word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>说明</w:t></w:r></w:p></w:document>',
                       "word/_rels/document.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="https://invalid.example/never-fetch" TargetMode="External" Type="hyperlink"/></Relationships>'})
        result = extractor.extract(raw, "docx")
        self.assertIn("external_links_not_followed", result["warnings"])
        self.assertTrue(result["truncated"])

    def test_cli_errors_never_echo_the_input_or_traceback(self):
        process = subprocess.run([sys.executable, "-I", str(Path(__file__).with_name("extractor.py")), "pdf"],
                                 input=b"invalid fixture document private-marker", capture_output=True, timeout=10)
        self.assertEqual(process.returncode, 2)
        self.assertEqual(process.stdout, b'{"error":"attachment_document_invalid"}')
        self.assertEqual(process.stderr, b"")

    def test_nested_textbox_paragraphs_are_not_counted_twice(self):
        raw = archive({"word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:pict><w:txbxContent><w:p><w:r><w:t>textbox bug</w:t></w:r></w:p></w:txbxContent></w:pict></w:r></w:p></w:document>'})
        result = extractor.extract(raw, "docx")
        self.assertEqual([r["text"] for r in result["records"]], ["textbox bug"])
        self.assertIn("unread_embedded_content", result["warnings"])

    def test_xlsx_hidden_sheets_are_not_omitted(self):
        workbook = Workbook()
        workbook.active["A1"] = "visible"
        hidden = workbook.create_sheet("hidden")
        hidden.sheet_state = "hidden"
        hidden["B2"] = "hidden bug"
        stream = io.BytesIO()
        workbook.save(stream)
        result = extractor.extract(stream.getvalue(), "xlsx")
        self.assertTrue(any(r["location"] == "hidden!B2" and r["text"] == "hidden bug" for r in result["records"]))


if __name__ == "__main__":
    unittest.main()
