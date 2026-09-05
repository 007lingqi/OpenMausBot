"""Untrusted document bytes in, bounded source-located JSON out. No file extraction or evaluation."""
import io
import json
import logging
import sys
import zipfile
from pathlib import PurePosixPath
from defusedxml import ElementTree
from openpyxl import load_workbook
from pypdf import PdfReader

logging.disable(logging.CRITICAL)
MAX_BYTES = 10 * 1024 * 1024
MAX_CHARACTERS = 250_000
MAX_RECORDS = 5000
WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def paragraph_text(paragraph):
    # Textbox paragraphs are emitted separately, not repeated through their outer paragraph.
    stack, output = [paragraph], []
    while stack:
        item = stack.pop()
        if item is not paragraph and item.tag == WORD + "p":
            continue
        if item.tag == WORD + "t":
            output.append(item.text or "")
        elif item.tag == WORD + "tab":
            output.append("\t")
        elif item.tag == WORD + "br":
            output.append("\n")
        stack.extend(reversed(list(item)))
    return "".join(output)


def checked_archive(raw):
    bundle = zipfile.ZipFile(io.BytesIO(raw))
    entries = bundle.infolist()
    if len(entries) > 2000 or sum(e.file_size for e in entries) > 32 * 1024 * 1024:
        raise ValueError("archive_limit")
    names = set()
    for entry in entries:
        name = entry.filename
        if name in names or name.startswith("/") or "\\" in name or ".." in PurePosixPath(name).parts or ":" in name:
            raise ValueError("archive_path")
        names.add(name)
        if entry.flag_bits & 1 or entry.file_size > 8 * 1024 * 1024 or entry.file_size > max(1, entry.compress_size) * 100:
            raise ValueError("archive_limit")
        if "vbaproject" in name.lower() or name.lower().endswith(".bin"):
            raise ValueError("active_content")
        if name.endswith((".xml", ".rels")):
            # defusedxml also rejects encoded DTD/entity payloads that a byte search misses.
            ElementTree.fromstring(bundle.read(entry), forbid_dtd=True, forbid_entities=True, forbid_external=True)
    return bundle


def extract(raw, format_name):
    if len(raw) > MAX_BYTES:
        raise ValueError("input_limit")
    records, warnings = [], set()
    characters = 0

    def add(location, value):
        nonlocal characters
        value = str(value)
        if not value.strip():
            return
        if characters + len(value) + len(location) > MAX_CHARACTERS or len(records) >= MAX_RECORDS:
            warnings.add("output_truncated")
            return
        records.append({"location": location, "text": value})
        characters += len(value) + len(location)

    if format_name in ("docx", "xlsx"):
        with checked_archive(raw) as bundle:
            for name in bundle.namelist():
                if any(part in name.lower() for part in ("/media/", "/embeddings/", "/charts/", "/drawings/", "comments")):
                    warnings.add("unread_embedded_content")
                if name.endswith(".rels"):
                    if any(item.get("TargetMode") == "External" for item in ElementTree.fromstring(bundle.read(name))):
                        warnings.add("external_links_not_followed")
            if format_name == "docx":
                names = ["word/document.xml"] + sorted(n for n in bundle.namelist() if n.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes")) and n.endswith(".xml"))
                for name in names:
                    root = ElementTree.fromstring(bundle.read(name))
                    parents = {child: parent for parent in root.iter() for child in parent}
                    labels = {}
                    for table_index, table in enumerate(root.iter(WORD + "tbl"), 1):
                        labels[table] = f"table:{table_index}"
                        for row_index, row in enumerate(table.findall(WORD + "tr"), 1):
                            labels[row] = f"row:{row_index}"
                            for cell_index, cell in enumerate(row.findall(WORD + "tc"), 1):
                                labels[cell] = f"cell:{cell_index}"
                    if any(item.tag in (WORD + "instrText", WORD + "fldSimple", WORD + "fldChar", WORD + "del", WORD + "ins", WORD + "altChunk", WORD + "object") for item in root.iter()):
                        warnings.add("unread_revision_or_field")
                    if any(item.tag in (WORD + "pict", WORD + "drawing") for item in root.iter()):
                        warnings.add("unread_embedded_content")
                    for index, paragraph in enumerate(root.iter(WORD + "p"), 1):
                        location, parent = [], parents.get(paragraph)
                        while parent is not None:
                            if parent in labels:
                                location.append(labels[parent])
                            parent = parents.get(parent)
                        prefix = ":".join([name, *reversed(location)])
                        add(f"{prefix}:paragraph:{index}", paragraph_text(paragraph))
            else:
                workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
                try:
                    if len(workbook.worksheets) > 50:
                        raise ValueError("sheet_limit")
                    for sheet in workbook.worksheets:
                        # Do not trust the optional dimension claimed by the producer.
                        sheet.reset_dimensions()
                        cells = 0
                        for row_index, row in enumerate(sheet.iter_rows(), 1):
                            if row_index > 20_000 or len(row) > 100:
                                raise ValueError("sheet_limit")
                            for cell in row:
                                cells += 1
                                if cells > 200_000:
                                    raise ValueError("sheet_limit")
                                if cell.value is not None:
                                    if cell.data_type == "f":
                                        warnings.add("formulas_not_evaluated")
                                    add(f"{sheet.title}!{cell.coordinate}", cell.value)
                finally:
                    workbook.close()
    elif format_name == "pdf":
        reader = PdfReader(io.BytesIO(raw), strict=True)
        if reader.is_encrypted or len(reader.pages) > 200:
            raise ValueError("pdf_encrypted_or_limit")
        for index, page in enumerate(reader.pages, 1):
            value = page.extract_text() or ""
            if not value.strip():
                warnings.add("unread_empty_or_scanned_page")
            if page.get("/Annots") or page.get("/Resources", {}).get("/XObject"):
                warnings.add("unread_images_or_annotations")
            add(f"page:{index}", value)
        root = reader.trailer["/Root"]
        if root.get("/AcroForm") or root.get("/Names") or root.get("/OpenAction"):
            warnings.add("unread_interactive_content")
    else:
        raise ValueError("format_unsupported")
    if not records:
        warnings.add("no_readable_text")
    return {"version": 1, "format": format_name, "records": records, "truncated": bool(warnings), "warnings": sorted(warnings)}


if __name__ == "__main__":
    try:
        result = extract(sys.stdin.buffer.read(MAX_BYTES + 1), sys.argv[1])
        sys.stdout.write(json.dumps(result, ensure_ascii=False))
    except Exception:
        # Never print parser errors, document contents, filenames or tracebacks.
        sys.stdout.write('{"error":"attachment_document_invalid"}')
        sys.exit(2)
